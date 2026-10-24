"""Preislistenimport der Lieferanten (Pflichtenheft Abschnitt 7): CSV/TXT, Excel, XML; Upload, FTP, SFTP oder URL."""
import csv
import ftplib
import io
import os
import urllib.error
import urllib.request
import xml.etree.ElementTree as ET
from decimal import Decimal, InvalidOperation

from datetime import timedelta

from django.db import transaction
from django.db.models import Q
from django.utils import timezone

from apps.stammdaten.models import Artikel, ArtikelLieferant, Preislistenartikel

from .models import Dateityp, Preislistenimport, Preislistenprofil, Quelle

FELDER = ("bestellnummer", "herstellernummer", "ean", "bezeichnung", "hersteller", "kategorie", "ek",
          "bestand", "lieferzeit", "gewicht")


class PreislistenFehler(Exception):
    pass


# ---- Abruf ------------------------------------------------------------------------------------------

def passwort(profil) -> str:
    return os.environ.get(profil.passwort_variable, "")


def abrufen(profil):
    """Holt die Datei vom FTP-/SFTP-Server bzw. der URL; das Passwort steht in der .env (PREISLISTE_PW_<ID>)."""
    if profil.quelle == Quelle.FTP:
        puffer = io.BytesIO()
        try:
            with ftplib.FTP(timeout=60) as ftp:
                ftp.connect(profil.adresse, profil.port or 21)
                ftp.login(profil.benutzer or "anonymous", passwort(profil))
                ftp.retrbinary(f"RETR {profil.pfad}", puffer.write)
        except (ftplib.all_errors, OSError) as fehler:
            raise PreislistenFehler(f"FTP-Abruf fehlgeschlagen: {fehler}")
        return os.path.basename(profil.pfad), puffer.getvalue()
    if profil.quelle == Quelle.SFTP:
        try:
            import paramiko
            transport = paramiko.Transport((profil.adresse, profil.port or 22))
            transport.connect(username=profil.benutzer, password=passwort(profil))
            sftp = paramiko.SFTPClient.from_transport(transport)
            with sftp.open(profil.pfad, "rb") as datei:
                inhalt = datei.read()
            sftp.close()
            transport.close()
        except Exception as fehler:
            raise PreislistenFehler(f"SFTP-Abruf fehlgeschlagen: {fehler}")
        return os.path.basename(profil.pfad), inhalt
    if profil.quelle == Quelle.URL:
        url = profil.url_adresse
        if not url.lower().startswith(("http://", "https://")):
            raise PreislistenFehler("Bitte die vollständige Adresse (https://…) im Feld „FTP-Server bzw. URL“ eintragen.")
        try:
            anfrage = urllib.request.Request(url)
            if profil.benutzer:
                if not passwort(profil):
                    raise PreislistenFehler(f"Das Passwort fehlt: {profil.passwort_variable} in der .env eintragen "
                                            "und den Dienst neu starten.")
                import base64
                anfrage.add_header("Authorization", "Basic " + base64.b64encode(
                    f"{profil.benutzer}:{passwort(profil)}".encode()).decode())
            with urllib.request.urlopen(anfrage, timeout=300) as antwort:
                return os.path.basename(url.split("?")[0].rstrip("/")) or "preisliste.csv", antwort.read()
        except PreislistenFehler:
            raise
        except urllib.error.HTTPError as fehler:
            hinweis = " – Benutzer oder Passwort prüfen." if fehler.code in (401, 403) else ""
            raise PreislistenFehler(f"Download fehlgeschlagen: Der Server antwortet mit {fehler.code} {fehler.reason}{hinweis}")
        except Exception as fehler:
            raise PreislistenFehler(f"Download fehlgeschlagen: {fehler}")
    raise PreislistenFehler("Dieses Profil wird per Datei-Upload gefüllt.")


# ---- Lesen ------------------------------------------------------------------------------------------

def spaltenindex(kopf, angabe: str, name: str):
    angabe = (angabe or "").strip()
    if not angabe:
        return None
    if angabe.isdigit():
        return int(angabe) - 1
    for i, k in enumerate(kopf or []):
        if str(k).strip().lower() == angabe.lower():
            return i
    raise PreislistenFehler(f"Spalte '{angabe}' ({name}) nicht in der Datei gefunden.")


def tabelle(profil, inhalt: bytes) -> list:
    """Zeilen als Listen (CSV, Excel)."""
    if profil.dateityp == Dateityp.CSV:
        try:
            text = inhalt.decode(profil.zeichensatz)
        except UnicodeDecodeError:
            raise PreislistenFehler(f"Die Datei lässt sich nicht als {profil.zeichensatz} lesen (Zeichensatz im Profil prüfen).")
        trenner = "\t" if profil.trennzeichen in ("\\t", "tab", "TAB") else (profil.trennzeichen or ";")
        kwargs = {"quotechar": profil.textbegrenzer} if profil.textbegrenzer else {"quoting": csv.QUOTE_NONE}
        return [z for z in csv.reader(io.StringIO(text), delimiter=trenner, **kwargs)]
    if profil.dateityp == Dateityp.XLSX:
        from openpyxl import load_workbook
        blatt = load_workbook(io.BytesIO(inhalt), read_only=True, data_only=True).worksheets[0]
        return [["" if v is None else v for v in z] for z in blatt.iter_rows(values_only=True)]
    if profil.dateityp == Dateityp.XLS:
        import xlrd
        blatt = xlrd.open_workbook(file_contents=inhalt, encoding_override="cp1252").sheet_by_index(0)
        return [blatt.row_values(i) for i in range(blatt.nrows)]
    raise PreislistenFehler("Unbekannter Dateityp.")


def text(wert) -> str:
    if isinstance(wert, float):
        return str(int(wert)) if wert == int(wert) else str(wert)
    return "" if wert is None else str(wert).strip()


def zeilen_lesen(profil, inhalt: bytes) -> list:
    """Liste von (zeilennummer, {feld: text}) gemäß Feldzuordnung."""
    zuordnung = {f: getattr(profil, f"sp_{f}") for f in FELDER}
    ergebnis = []
    if profil.dateityp == Dateityp.XML:
        if not profil.xml_satz:
            raise PreislistenFehler("Für XML bitte das Element je Artikel im Profil angeben.")
        try:
            wurzel = ET.fromstring(inhalt)
        except ET.ParseError as fehler:
            raise PreislistenFehler(f"XML nicht lesbar: {fehler}")
        for nr, satz in enumerate((e for e in wurzel.iter() if e.tag.rsplit("}", 1)[-1] == profil.xml_satz), start=1):
            werte = {k.rsplit("}", 1)[-1]: (v or "") for k, v in satz.attrib.items()}
            werte.update({c.tag.rsplit("}", 1)[-1]: (c.text or "").strip() for c in satz})
            ergebnis.append((nr, {f: text(werte.get(a, "")) for f, a in zuordnung.items() if a}))
        return ergebnis
    zeilen = tabelle(profil, inhalt)
    kopf = zeilen[0] if profil.kopfzeile and zeilen else None
    index = {f: spaltenindex(kopf, a, f) for f, a in zuordnung.items()}
    if index["bestellnummer"] is None or index["ek"] is None:
        raise PreislistenFehler("Bestellnummer und EK müssen zugeordnet sein.")
    for nr, z in enumerate(zeilen, start=1):
        if nr < max(profil.erste_datenzeile, 2 if profil.kopfzeile else 1) or not any(text(c) for c in z):
            continue
        ergebnis.append((nr, {f: text(z[i]) if i is not None and i < len(z) else "" for f, i in index.items()}))
    return ergebnis


def zahl(wert: str, dezimal: str) -> Decimal:
    t = wert.replace("€", "").replace("EUR", "").replace(" ", "").replace("\xa0", "")
    t = t.replace(".", "").replace(",", ".") if dezimal == "," and ("," in t or t.count(".") > 1) else t.replace(",", "") if dezimal == "." else t
    try:
        return Decimal(t)
    except InvalidOperation:
        raise ValueError(f"Zahl nicht lesbar: '{wert}'")


# ---- Import -----------------------------------------------------------------------------------------

def artikel_finden(zeile: dict):
    """Eigener Artikel (Typ A/L/G) zur Zeile: über die Herstellernummer, ersatzweise über die EAN."""
    if zeile.get("herstellernummer"):
        a = Artikel.objects.filter(herstellernummer__iexact=zeile["herstellernummer"]).first()
        if a:
            return a
    if zeile.get("ean"):
        return Artikel.objects.filter(ean=zeile["ean"]).first()
    return None


def vorschau(profil, inhalt: bytes, anzahl: int = 15) -> dict:
    """Erste Zeilen mit Zuordnung und Abgleichergebnis anzeigen, ohne etwas zu speichern."""
    zeilen = zeilen_lesen(profil, inhalt)
    zeilen_vorschau = []
    for nr, z in zeilen[:anzahl]:
        bekannt = Preislistenartikel.objects.filter(lieferant=profil.lieferant, bestellnummer=z.get("bestellnummer", "")).first()
        artikel = artikel_finden(z)
        try:
            ek = zahl(z.get("ek", ""), profil.dezimaltrennzeichen)
        except ValueError as fehler:
            ek = str(fehler)
        if bekannt:
            ergebnis = "in der Preisliste vorhanden – wird aktualisiert"
        elif artikel or profil.neue_artikel:
            ergebnis = "neu in der Preisliste"
        else:
            ergebnis = "wird übersprungen (kein eigener Artikel)"
        zeilen_vorschau.append({"nr": nr, "daten": z, "ek": ek, "artikel": artikel, "ergebnis": ergebnis})
    return {"gesamt": len(zeilen), "zeilen": zeilen_vorschau}


SPERRE_STUNDEN = 3  # ein abgestürzter Lauf blockiert das Profil nicht länger als so lange


def sperre_holen(profil) -> None:
    """Verhindert, dass dieselbe Preisliste zweimal gleichzeitig importiert wird (Cron und Handstart)."""
    jetzt = timezone.now()
    erfolgt = Preislistenprofil.objects.filter(pk=profil.pk).filter(
        Q(laeuft_seit__isnull=True) | Q(laeuft_seit__lt=jetzt - timedelta(hours=SPERRE_STUNDEN))).update(laeuft_seit=jetzt)
    if not erfolgt:
        raise PreislistenFehler("Für diese Preisliste läuft bereits ein Import – bitte warten, bis er fertig ist.")


def sperre_freigeben(profil) -> None:
    Preislistenprofil.objects.filter(pk=profil.pk).update(laeuft_seit=None)


def importieren(profil, inhalt: bytes, dateiname: str = "", benutzer=None) -> Preislistenimport:
    """Liest die Preisliste in die Tabelle der Preislistenartikel ein. Bestehende Zeilen werden einmal in den Speicher geladen;
    geschrieben wird nur, was neu ist oder sich ändert (in Blöcken). Passt eine Zeile zu einem eigenen Artikel
    (Herstellernummer, ersatzweise EAN), wird sie mit ihm verknüpft und der Lieferantenpreis am Artikel mitgeführt."""
    zeilen = zeilen_lesen(profil, inhalt)  # Formatfehler vor dem Sperren melden
    sperre_holen(profil)
    try:
        return _importieren(profil, zeilen, dateiname, benutzer)
    finally:
        sperre_freigeben(profil)


def _importieren(profil, zeilen: list, dateiname: str, benutzer) -> Preislistenimport:
    protokoll = Preislistenimport(profil=profil, dateiname=dateiname[:200], gelesen=len(zeilen), benutzer=benutzer)
    fehler, jetzt, lieferant = [], timezone.now(), profil.lieferant

    # Bestehende Preislistenzeilen: Bestellnummer -> Zeilenwerte (kompakt, ohne volle Modellobjekte)
    vorhanden = {}
    for pk, bn, ek, bestand, lz, artikel_id, bez, hersteller in Preislistenartikel.objects.filter(lieferant=lieferant).values_list(
            "id", "bestellnummer", "ek", "bestand", "lieferzeit_tage", "artikel_id", "bezeichnung", "hersteller"):
        vorhanden[bn] = {"id": pk, "ek": ek, "bestand": bestand, "lz": lz, "artikel_id": artikel_id, "bez": bez, "hersteller": hersteller}
    nach_hnr, nach_ean = {}, {}
    for artikel_id, hnr, ean in Artikel.objects.values_list("id", "herstellernummer", "ean"):
        if hnr:
            nach_hnr.setdefault(hnr.lower(), artikel_id)
        if ean:
            nach_ean.setdefault(ean, artikel_id)
    lieferantenpreise = {al.bestellnummer: al for al in ArtikelLieferant.objects.filter(lieferant=lieferant)}

    neue, geaenderte = {}, {}  # Bestellnummer -> Preislistenartikel (noch nicht bzw. neu zu schreiben)
    al_neu, al_geaendert = [], []

    def schreiben():
        if neue:
            Preislistenartikel.objects.bulk_create(list(neue.values()), batch_size=1000)
            neue.clear()
        if geaenderte:
            Preislistenartikel.objects.bulk_update(
                list(geaenderte.values()), ["ek", "bestand", "lieferzeit_tage", "artikel", "bezeichnung", "hersteller",
                                            "preisliste", "aktualisiert_am"], batch_size=1000)
            geaenderte.clear()
        if al_neu:
            ArtikelLieferant.objects.bulk_create(al_neu, batch_size=500)
            al_neu.clear()
        if al_geaendert:
            ArtikelLieferant.objects.bulk_update(al_geaendert, ["ek", "bestand", "lieferzeit_tage", "preisliste", "aktualisiert_am"], batch_size=500)
            al_geaendert.clear()

    for nr, z in zeilen:
        try:
            bestellnr = z.get("bestellnummer", "").strip()[:100]
            if not bestellnr:
                raise ValueError("Bestellnummer fehlt.")
            ek = zahl(z.get("ek", ""), profil.dezimaltrennzeichen).quantize(Decimal("0.0001"))
            if ek < 0:
                raise ValueError("EK ist negativ.")
            bestand = int(zahl(z["bestand"], profil.dezimaltrennzeichen)) if z.get("bestand") else None
            lieferzeit = int(zahl(z["lieferzeit"], profil.dezimaltrennzeichen)) if z.get("lieferzeit") else None
            hnr, ean = z.get("herstellernummer", "").strip()[:100], z.get("ean", "").strip()[:14]
            bezeichnung, hersteller = z.get("bezeichnung", "").strip()[:255], z.get("hersteller", "").strip()[:150]

            artikel_id = nach_hnr.get(hnr.lower()) if hnr else None
            if artikel_id is None and ean:
                artikel_id = nach_ean.get(ean)
            alt = vorhanden.get(bestellnr)

            if alt is None:
                if not profil.neue_artikel and artikel_id is None:
                    protokoll.uebersprungen += 1
                    continue
                if not bezeichnung:
                    raise ValueError("Zeile ohne Bezeichnung kann nicht übernommen werden.")
                gewicht = zahl(z["gewicht"], profil.dezimaltrennzeichen) if z.get("gewicht") else None
                eintrag = Preislistenartikel(
                    lieferant=lieferant, bestellnummer=bestellnr, herstellernummer=hnr, ean=ean, bezeichnung=bezeichnung,
                    hersteller=hersteller, kategorie=z.get("kategorie", "").strip()[:150], ek=ek, bestand=bestand,
                    lieferzeit_tage=lieferzeit, gewicht_kg=gewicht, artikel_id=artikel_id, preisliste=profil.name[:150],
                    aktualisiert_am=jetzt)
                neue[bestellnr] = eintrag
                vorhanden[bestellnr] = {"id": None, "ek": ek, "bestand": bestand, "lz": lieferzeit, "artikel_id": artikel_id,
                                        "bez": bezeichnung, "hersteller": hersteller, "objekt": eintrag}
                protokoll.neu += 1
            else:
                neu_bestand = bestand if bestand is not None else alt["bestand"]
                neu_lz = lieferzeit if lieferzeit is not None else alt["lz"]
                neu_bez = bezeichnung or alt["bez"]
                neu_artikel = artikel_id if artikel_id is not None else alt["artikel_id"]
                aenderung = (alt["ek"] != ek or alt["bestand"] != neu_bestand or alt["lz"] != neu_lz or alt["bez"] != neu_bez
                             or alt["hersteller"] != (hersteller or alt["hersteller"]) or alt["artikel_id"] != neu_artikel)
                if aenderung:
                    objekt = alt.get("objekt")
                    if objekt is not None:  # doppelte Zeile in derselben Datei: noch nicht geschriebenen Eintrag anpassen
                        objekt.ek, objekt.bestand, objekt.lieferzeit_tage, objekt.bezeichnung = ek, neu_bestand, neu_lz, neu_bez
                        objekt.artikel_id = neu_artikel
                    else:
                        geaenderte[bestellnr] = Preislistenartikel(
                            id=alt["id"], lieferant=lieferant, bestellnummer=bestellnr, ek=ek, bestand=neu_bestand,
                            lieferzeit_tage=neu_lz, artikel_id=neu_artikel, bezeichnung=neu_bez,
                            hersteller=hersteller or alt["hersteller"], preisliste=profil.name[:150], aktualisiert_am=jetzt)
                    alt.update(ek=ek, bestand=neu_bestand, lz=neu_lz, bez=neu_bez, artikel_id=neu_artikel,
                               hersteller=hersteller or alt["hersteller"])
                    protokoll.geaendert += 1
                else:
                    protokoll.unveraendert += 1
                artikel_id = neu_artikel

            if artikel_id is not None:  # Lieferantenpreis am eigenen Artikel mitführen
                al = lieferantenpreise.get(bestellnr)
                if al is None:
                    al = ArtikelLieferant(artikel_id=artikel_id, lieferant=lieferant, bestellnummer=bestellnr, ek=ek, bestand=bestand,
                                          lieferzeit_tage=lieferzeit, preisliste=profil.name[:150], aktualisiert_am=jetzt)
                    lieferantenpreise[bestellnr] = al
                    al_neu.append(al)
                elif al.ek != ek or (bestand is not None and al.bestand != bestand) or al.artikel_id != artikel_id:
                    al.ek, al.artikel_id, al.preisliste, al.aktualisiert_am = ek, artikel_id, profil.name[:150], jetzt
                    if bestand is not None:
                        al.bestand = bestand
                    if lieferzeit is not None:
                        al.lieferzeit_tage = lieferzeit
                    if al.pk:
                        al_geaendert.append(al)
            if len(neue) >= 1000 or len(geaenderte) >= 1000:
                schreiben()
        except (ValueError, PreislistenFehler) as grund:
            protokoll.fehler += 1
            if len(fehler) < 200:
                fehler.append({"zeile": nr, "grund": str(grund), "daten": z})
    schreiben()
    protokoll.fehlerzeilen = fehler
    protokoll.meldung = f"{protokoll.neu} neu, {protokoll.geaendert} geändert, {protokoll.unveraendert} unverändert, " \
                        f"{protokoll.uebersprungen} übersprungen, {protokoll.fehler} Fehler"
    protokoll.save()
    Preislistenprofil.objects.filter(pk=profil.pk).update(letzter_lauf=jetzt)
    profil.letzter_lauf = jetzt
    return protokoll


def faellige_profile(jetzt=None):
    jetzt = jetzt or timezone.now()
    for p in Preislistenprofil.objects.filter(aktiv=True, intervall_stunden__gt=0).exclude(quelle=Quelle.UPLOAD):
        if p.letzter_lauf is None or (jetzt - p.letzter_lauf).total_seconds() >= p.intervall_stunden * 3600:
            yield p
